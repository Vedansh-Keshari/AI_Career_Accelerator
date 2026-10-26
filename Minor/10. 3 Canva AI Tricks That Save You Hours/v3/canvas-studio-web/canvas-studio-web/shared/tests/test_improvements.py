"""Tests for the senior-developer improvements. Run:  python -m pytest shared/tests -q"""
import io
import json
import os
import sys
import urllib.parse
from pathlib import Path

import pytest
import requests
from PIL import Image, ImageChops, ImageStat

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from studio import ai as ai_mod  # noqa: E402
from studio import elements as E  # noqa: E402
from studio import generator as G  # noqa: E402
from studio import models as M  # noqa: E402
from studio import office  # noqa: E402
from studio import options as O  # noqa: E402
from studio import render as R  # noqa: E402
from studio.canva_client import CanvaClient, CanvaError, make_pkce, push_project  # noqa: E402
from studio.mock_canva import MockCanva  # noqa: E402
from studio.paths import load_config  # noqa: E402
from studio.secrets_store import SecretStore, Settings  # noqa: E402

GOLDEN = Path(__file__).parent / "golden"


# ------------------------------------------------------------------ editable Office export
def test_pptx_has_native_editable_objects(tmp_path):
    from pptx import Presentation
    p = G.project_from_example("yoga-presentation")
    photo = tmp_path / "pic.png"
    Image.new("RGB", (300, 200), "#ff0066").save(photo)
    p["slides"][2]["image"] = str(photo)
    p["slides"][1]["free"] = True
    p["slides"][1]["elements"] = E.default_elements(p, 1) + [E.blank_element("rect"), E.blank_element("ellipse")]
    prs = Presentation(io.BytesIO(office.presentation_to_pptx(p)))
    assert len(prs.slides) == 6
    first = [s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text]
    assert first[0].text_frame.text == "Benefits of Yoga"        # real text, not a picture of text
    kinds = {s.shape_type for s in prs.slides[1].shapes}
    assert len(prs.slides[1].shapes) >= 5 and 1 in {int(k) for k in kinds}   # autoshapes present
    assert any(s.shape_type == 13 for s in prs.slides[2].shapes)             # the user's photo is a picture object
    bullets = [s.text_frame.text for s in prs.slides[1].shapes if s.has_text_frame and "\u2022" in s.text_frame.text]
    assert bullets and bullets[0].startswith("\u2022 ")
    for fmt, size in (("Presentation (4:3)", (10.0, 7.5)), ("Presentation (Mobile-First)", (7.5, 13.333))):
        p["options"]["format"] = fmt
        prs = Presentation(io.BytesIO(office.presentation_to_pptx(p)))
        assert prs.slide_width / 914400 == pytest.approx(size[0], abs=0.01) and prs.slide_height / 914400 == pytest.approx(size[1], abs=0.01)


def test_pptx_has_no_notes_slides_so_packaged_apps_need_no_notes_template():
    from pptx import Presentation
    prs = Presentation(io.BytesIO(office.presentation_to_pptx(G.project_from_example("yoga-presentation"))))
    assert not any(s.has_notes_slide for s in prs.slides)


def test_docx_has_headings_and_bullets():
    from docx import Document
    p = G.project_from_example("workout-blog-refined")
    doc = Document(io.BytesIO(office.document_to_docx(p)))
    styles = [x.style.name for x in doc.paragraphs]
    assert styles[0] == "Title" and styles.count("Heading 1") == len(p["doc"]["sections"])
    assert "List Bullet" in styles                                        # the "Key takeaways" list
    assert doc.core_properties.title == p["title"]


def test_push_falls_back_to_pdf_when_canva_refuses_office_files(canva):
    client, mock = canva
    signed_in(client)
    mock.reject_office = True
    result = push_project(client, G.project_from_example("yoga-presentation"))
    assert result["editable"] is False and "images" in result["kind"] and "did not accept" in result["note"]
    assert mock.files[mock.designs[result["id"]]["file"]][1] == "application/pdf"


def test_push_falls_back_when_office_libraries_missing(canva, monkeypatch):
    client, mock = canva
    signed_in(client)

    def boom(_):
        raise ImportError("no python-pptx")
    monkeypatch.setattr(office, "presentation_to_pptx", boom)
    result = push_project(client, G.project_from_example("yoga-presentation"))
    assert result["editable"] is False and "not installed" in result["note"]


def test_push_reports_progress(canva):
    client, _ = canva
    signed_in(client)
    seen = []
    push_project(client, G.project_from_example("yoga-presentation"), lambda stage, step=0: seen.append(stage))
    assert seen[0] == "Preparing your pages" and "Uploading to Canva" in seen


# ------------------------------------------------------------------ retries
def test_get_requests_retry_temporary_failures(canva):
    client, mock = canva
    signed_in(client)
    mock.flaky = [503, 502]
    assert client.profile()["display_name"] == "vedansh"       # succeeded on the third try
    mock.flaky = [503, 503, 503, 503]
    with pytest.raises(CanvaError):
        client.list_designs()


def test_post_is_not_repeated_after_server_error_but_429_is(canva):
    client, mock = canva
    signed_in(client)
    mock.flaky = [500]
    with pytest.raises(CanvaError):
        client.create_design("t", "Mnone", 100, 100)
    assert len(mock.designs) == 0 and mock.flaky == []          # exactly one attempt was made
    mock.flaky = [429]
    with pytest.raises(CanvaError, match="asset not found"):    # retried once after the 429, then real answer
        client.create_design("t", "Mnone", 100, 100)


# ------------------------------------------------------------------ secrets
def test_secret_store_encrypts_and_roundtrips(tmp_path):
    s = SecretStore(tmp_path, use_keyring=False)
    s.set("tokens_web", json.dumps({"access_token": "SUPER-SECRET-123"}))
    assert b"SUPER-SECRET-123" not in (tmp_path / "secrets.enc").read_bytes()
    assert SecretStore(tmp_path, use_keyring=False).get("tokens_web") and "SUPER" in SecretStore(tmp_path, use_keyring=False).get("tokens_web")
    s.delete("tokens_web")
    assert s.get("tokens_web") is None


def test_secret_store_survives_a_damaged_file(tmp_path):
    s = SecretStore(tmp_path, use_keyring=False)
    s.set("a", "1")
    (tmp_path / "secrets.enc").write_bytes(b"garbage")
    assert s.get("a") is None                                   # no crash, just empty
    s.set("a", "2")
    assert s.get("a") == "2"


def test_secret_store_uses_os_keyring_when_available(tmp_path):
    import keyring
    from keyring.backend import KeyringBackend

    class Memory(KeyringBackend):
        priority = 1
        def __init__(self):
            self.data = {}
        def get_password(self, service, username):
            return self.data.get((service, username))
        def set_password(self, service, username, password):
            self.data[(service, username)] = password
        def delete_password(self, service, username):
            self.data.pop((service, username), None)

    old = keyring.get_keyring()
    mem = Memory()
    keyring.set_keyring(mem)
    try:
        s = SecretStore(tmp_path, use_keyring=True)
        assert s.protected_by == "keyring" and not (tmp_path / "secrets.key").exists()
        assert any(k[1] == "master-key" for k in mem.data)
        s.set("x", "y")
        assert SecretStore(tmp_path, use_keyring=True).get("x") == "y"   # same key fetched back from the keyring
    finally:
        keyring.set_keyring(old)


def test_secret_store_falls_back_when_keyring_is_broken(tmp_path, monkeypatch):
    import keyring

    def broken(*a, **k):
        raise keyring.errors.NoKeyringError("none")
    monkeypatch.setattr(keyring, "get_password", broken)
    s = SecretStore(tmp_path, use_keyring=True)
    assert s.protected_by == "file" and (tmp_path / "secrets.key").exists()
    s.set("k", "v")
    assert s.get("k") == "v"


def test_tokens_are_never_written_in_plain_text(canva):
    client, _ = canva
    signed_in(client)
    token = client.tokens.load()["access_token"]
    for f in Path(client.tokens.secrets.home).iterdir():
        if f.is_file():
            assert token.encode() not in f.read_bytes()


def test_config_prefers_saved_settings_over_file(tmp_path, monkeypatch):
    monkeypatch.delenv("CANVA_CLIENT_ID", raising=False)
    monkeypatch.delenv("CANVA_CLIENT_SECRET", raising=False)
    secrets, settings = SecretStore(tmp_path, use_keyring=False), Settings(tmp_path)
    assert load_config("web", secrets, settings)["demo"] is True
    settings.update({"canva": {"client_id": "OC-mine"}})
    secrets.set("canva_client_secret", "shh")
    cfg = load_config("web", secrets, settings)
    assert cfg["demo"] is False and cfg["client_id"] == "OC-mine" and cfg["client_secret"] == "shh"
    secrets.delete("canva_client_secret")
    with pytest.raises(SystemExit):
        load_config("web", secrets, settings)                    # id without secret is a clear error, not a silent failure


# ------------------------------------------------------------------ project format
def test_old_projects_are_migrated_without_losing_content():
    v1 = {"id": "p_0123456789", "type": "presentation", "title": "Old", "options": {"style": "Bold"},
          "slides": [{"title": "A", "body": "b", "layout": "auto", "image": ""}], "tune": {"boldness": 2, "text_scale": 1.0}}
    p = M.clean_project(v1)
    assert p["schema_version"] == M.SCHEMA_VERSION and p["slides"][0]["free"] is False and p["slides"][0]["elements"] == []
    assert p["tune"]["font"] == "auto" and p["tune"]["boldness"] == 2 and p["options"]["style"] == "Bold"
    assert R.render_pages(p, 0.2)


def test_revision_conflicts_are_detected(tmp_path):
    store = M.Store(tmp_path)
    p = store.save(G.project_from_example("cafe-logo"))
    assert p["rev"] == 1
    p2 = store.save(p, expected_rev=1)
    assert p2["rev"] == 2
    with pytest.raises(M.Conflict):
        store.save(p, expected_rev=1)                            # a second tab still holding rev 1
    assert store.save(p2, bump=False)["rev"] == 2                # server-owned updates keep the rev


def test_save_retries_when_windows_briefly_locks_the_file(tmp_path, monkeypatch):
    store = M.Store(tmp_path)
    p = G.project_from_example("cafe-logo")
    real, calls = os.replace, {"n": 0}

    def flaky(src, dst):
        calls["n"] += 1
        if calls["n"] <= 2:
            raise PermissionError(13, "The process cannot access the file because it is being used by another process")
        return real(src, dst)
    monkeypatch.setattr(os, "replace", flaky)
    assert store.save(p)["id"] == p["id"] and calls["n"] == 3


def test_upload_sandbox_and_cleanup(tmp_path):
    store = M.Store(tmp_path)
    mine = Path(store.save_upload(b"\xff\xd8x", "jpg"))
    assert store.is_upload(str(mine)) and not store.is_upload("/etc/passwd") and not store.is_upload("")
    assert not store.is_upload(str(tmp_path / "uploads" / ".." / "projects" / "x.json"))
    old_unused = Path(store.save_upload(b"a", "png"))
    used = Path(store.save_upload(b"b", "png"))
    p = G.project_from_example("yoga-presentation")
    p["slides"][0]["image"] = str(used)
    store.save(p)
    for f in (old_unused, used):
        os.utime(f, (0, 0))                                      # make both "old"
    assert store.cleanup_uploads(max_age_days=1) == 1
    assert used.exists() and not old_unused.exists() and mine.exists()


def test_windows_reserved_names_never_become_file_names():
    for name in ("CON", "nul", "Com1", "lpt9", "aux"):
        assert M.slug(name).startswith("design-")
    assert M.slug("") == "design" and M.slug("***") == "design" and len(M.slug("x" * 300)) <= 60


# ------------------------------------------------------------------ free-layout editor engine
@pytest.mark.parametrize("style", O.STYLES)
def test_default_elements_match_every_style(style):
    p = G.project_from_example("yoga-presentation")
    p["options"]["style"] = style
    for i in range(len(p["slides"])):
        els = E.default_elements(p, i)
        assert els and all(0 <= e["x"] <= 1 and 0 < e["w"] <= 1.4 for e in els)
        p["slides"][i]["elements"], p["slides"][i]["free"] = els, True
    assert len(R.render_pages(p, 0.15)) == 6


def test_every_element_type_renders_and_bad_input_is_repaired(tmp_path):
    p = G.project_from_example("yoga-presentation")
    pic = tmp_path / "p.png"
    Image.new("RGB", (200, 100), "#00ff00").save(pic)
    p["slides"][0]["free"] = True
    p["slides"][0]["elements"] = [
        {"type": "rect", "x": 0.0, "y": 0.0, "w": 1.0, "h": 1.0, "fill": "#000000"},
        {"type": "ellipse", "x": 0.1, "y": 0.1, "w": 0.3, "h": 0.3, "fill": "accent", "opacity": 0.5},
        {"type": "text", "x": 0.5, "y": 0.5, "w": 0.4, "h": 0.2, "text": "Hello\n- one\n- two " * 30, "size": 9999, "color": "not-a-colour", "align": "sideways"},
        {"type": "image", "x": 0.6, "y": 0.1, "w": 0.3, "h": 0.3, "image": str(pic)},
        {"type": "image", "x": 0.0, "y": 0.6, "w": 0.2, "h": 0.2, "image": "/missing.png"},
        {"type": "banana"}, "junk", None,
    ]
    clean = M.clean_project(p)["slides"][0]["elements"]
    assert [e["type"] for e in clean] == ["rect", "ellipse", "text", "image", "image"]
    assert clean[2]["size"] == 400 and clean[2]["align"] == "left" and clean[2]["color"] == "text"
    img = R.render_slide(M.clean_project(p), 0, size=(480, 270))
    assert img.getpixel((5, 5)) == (0, 0, 0)                       # the black rectangle covers the slide
    assert img.getpixel((int(480 * 0.75), int(270 * 0.25)))[1] > 200   # the green photo is visible


def test_text_shrinks_to_fit_its_box():
    p = G.project_from_example("yoga-presentation")
    look = R.Look(p)
    el = M.clean_element({"type": "text", "x": 0, "y": 0, "w": 0.3, "h": 0.1, "text": "word " * 80, "size": 120})
    fit = R.fit_element_text(el, look, 0.3 * 1920, 0.1 * 1080, 1.0)
    assert fit["size"] < 60 and fit["height"] <= 0.1 * 1080 + 1


def test_new_fonts_and_brand_palette_render():
    p = G.project_from_example("yoga-presentation")
    for font in O.FONTS.values():
        p["tune"]["font"] = font
        assert R.render_pages(M.clean_project(p), 0.15)
    p["brand"] = {"primary": "#7a1f5c", "accent": "#f4a261", "name": "Acme"}
    p["options"]["color"] = "Brand"
    clean = M.clean_project(p)
    assert clean["options"]["color"] == "Brand" and R.Look(clean).pal["primary"] == "#7a1f5c"
    bad = M.clean_project({**p, "brand": {"primary": "red", "accent": "#fff"}})
    assert "brand" not in bad and bad["options"]["color"] != "Brand"


# ------------------------------------------------------------------ AI starter text
class FakeSession:
    def __init__(self, status=200, text="", boom=False):
        self.status, self.text, self.boom, self.calls = status, text, boom, []

    def post(self, url, **kw):
        self.calls.append((url, kw))
        if self.boom:
            raise requests.ConnectionError("offline")
        r = requests.Response()
        r.status_code = self.status
        r._content = json.dumps({"content": [{"type": "text", "text": self.text}]}).encode()
        return r


GOOD_DECK = json.dumps({"title": "Bees", "slides": [{"title": "Bees", "body": "A buzzing intro"}, {"title": "Why", "body": "- honey\n- pollination"}, {"title": "Join", "body": "Start today"}]})


def test_ai_outline_is_used_and_validated():
    s = FakeSession(text="```json\n" + GOOD_DECK + "\n```")
    got = ai_mod.generate("presentation", "bees", {}, "Blog post", "sk-test", session=s)
    assert got["title"] == "Bees" and len(got["slides"]) == 3
    url, kw = s.calls[0]
    assert url.endswith("/v1/messages") and kw["headers"]["x-api-key"] == "sk-test" and "<idea>" in kw["json"]["messages"][0]["content"]
    project = G.build_project("presentation", "bees", {}, ai=lambda *a: ai_mod.generate(*a, "sk-test", session=s))
    assert project["generated_by"] == "ai" and project["slides"][0]["layout"] == "title" and project["slides"][-1]["layout"] == "cta"


@pytest.mark.parametrize("session,expect", [
    (FakeSession(text="I cannot do JSON"), "did not return JSON"),
    (FakeSession(text="{not json}"), "could not be read"),
    (FakeSession(text='{"title":"x","slides":[{"title":"only one","body":""}]}'), "too short"),
    (FakeSession(status=401), "rejected"),
    (FakeSession(status=429), "busy"),
    (FakeSession(status=500), "HTTP 500"),
    (FakeSession(boom=True), "Could not reach"),
])
def test_ai_failures_fall_back_to_templates_with_a_reason(session, expect):
    notices = []
    project = G.build_project("presentation", "about tea", {}, notices=notices, ai=lambda *a: ai_mod.generate(*a, "sk-x", session=session))
    assert project["generated_by"] == "templates" and len(project["slides"]) >= 3
    assert len(notices) == 1 and expect in notices[0] and "templates" in notices[0]


def test_ai_documents_logos_and_clipping():
    s = FakeSession(text=json.dumps({"title": "T" * 500, "sections": [{"heading": "H", "body": "b" * 9000}] * 30}))
    doc = ai_mod.generate("document", "x", {}, "Blog post", "k", session=s)
    assert len(doc["title"]) == 120 and len(doc["sections"]) == 12 and len(doc["sections"][0]["body"]) == 3000
    s = FakeSession(text='{"brand": "Bee Co", "tagline": "Sweet work"}')
    p = G.build_project("logo", "a honey shop", {}, ai=lambda *a: ai_mod.generate(*a, "k", session=s))
    assert p["logo"]["brand"] == "Bee Co" and p["title"] == "Bee Co"
    with pytest.raises(ai_mod.AIError):
        ai_mod.generate("logo", "x", {}, "Blog post", "", session=s)       # no key


# ------------------------------------------------------------------ golden images (catch accidental layout changes)
def _cases():
    yoga = G.project_from_example("yoga-presentation")
    free = G.project_from_example("yoga-presentation")
    free["slides"][2]["elements"], free["slides"][2]["free"] = E.default_elements(free, 2), True
    neo = G.project_from_example("yoga-presentation")
    neo["options"].update(style="NeoDigital", color="Purple")
    return {
        "slide_title_organic": lambda: R.render_slide(yoga, 0, size=(480, 270)),
        "slide_bullets_organic": lambda: R.render_slide(yoga, 2, size=(480, 270)),
        "slide_free_bullets": lambda: R.render_slide(free, 2, size=(480, 270)),
        "slide_neodigital": lambda: R.render_slide(neo, 1, size=(480, 270)),
        "doc_page1": lambda: R.render_pages(G.project_from_example("workout-blog-refined"), 0.3)[0],
        "logo_badge": lambda: R.render_logo(G.project_from_example("cafe-logo"), 0, 0.3),
        "logo_bright_stack": lambda: R.render_logo(G.project_from_example("cafe-logo-bright"), 2, 0.3),
    }


@pytest.mark.parametrize("name", list(_cases()))
def test_golden_image(name):
    """Compares against stored pictures. Small differences (font smoothing on another OS) are tolerated;
    a real layout change is not. To accept an intended change: UPDATE_GOLDEN=1 python -m pytest -k golden"""
    img = _cases()[name]().convert("RGB")
    path = GOLDEN / f"{name}.png"
    if os.environ.get("UPDATE_GOLDEN") or not path.exists():
        GOLDEN.mkdir(exist_ok=True)
        img.save(path)
        if not os.environ.get("UPDATE_GOLDEN"):
            pytest.skip("baseline created")
    base = Image.open(path).convert("RGB")
    assert base.size == img.size
    mean = sum(ImageStat.Stat(ImageChops.difference(base, img)).mean) / 3
    assert mean < 3.0, f"{name} differs from its baseline (mean pixel difference {mean:.2f})"


# ------------------------------------------------------------------ fixtures shared with test_shared
@pytest.fixture()
def canva(tmp_path):
    mock = MockCanva(8796).start()
    cfg = load_config("web")
    cfg.update(authorize_url=f"{mock.base}/api/oauth/authorize", token_url=f"{mock.base}/rest/v1/oauth/token",
               api_base=f"{mock.base}/rest/v1", redirect_uri="http://127.0.0.1:5000/oauth/callback")
    client = CanvaClient(cfg, home=tmp_path, sleep=lambda s: None, secret_store=SecretStore(tmp_path, use_keyring=False))
    yield client, mock
    mock.stop()


def signed_in(client):
    verifier, challenge = make_pkce()
    url = client.authorize_url("st", challenge)
    q = dict(urllib.parse.parse_qsl(urllib.parse.urlparse(url).query))
    form = {k: q[k] for k in ("client_id", "redirect_uri", "state", "code_challenge", "scope")}
    form.update(email="vedansh@example.com", action="allow")
    back = requests.post(url.split("?")[0], data=form, allow_redirects=False).headers["Location"]
    code = dict(urllib.parse.parse_qsl(urllib.parse.urlparse(back).query))["code"]
    client.exchange_code(code, verifier)
