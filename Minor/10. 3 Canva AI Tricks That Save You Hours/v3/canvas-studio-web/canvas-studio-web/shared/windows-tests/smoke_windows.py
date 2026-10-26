"""Environment and Windows-specific smoke checks for Canvas Studio.

Run:  python smoke_windows.py [--json results.json]
Works on any OS (Windows-only checks are marked SKIP elsewhere). Exit code 1 if anything FAILS.
PASS = fine, WARN = works but worth a look, FAIL = must be fixed, SKIP = not applicable here.
"""
from __future__ import annotations

import importlib
import json
import os
import platform
import socket
import sys
import tempfile
import threading
import time
import urllib.request
from pathlib import Path

HERE = Path(__file__).resolve().parent
SHARED = HERE.parent
sys.path.insert(0, str(SHARED))
os.environ.setdefault("CANVAS_NO_KEYRING", "")
IS_WIN = sys.platform == "win32"
results: list[dict] = []


def record(name: str, status: str, detail: str = "") -> None:
    results.append({"check": name, "status": status, "detail": detail})
    print(f"  [{status:<4}] {name}" + (f" - {detail}" if detail else ""))


def check(name: str):
    def wrap(fn):
        try:
            out = fn()
            if isinstance(out, tuple):
                record(name, *out)
            else:
                record(name, "PASS", out or "")
        except Exception as exc:  # noqa: BLE001 - a smoke test reports, it never crashes
            record(name, "FAIL", f"{type(exc).__name__}: {exc}")
        return fn
    return wrap


print(f"Canvas Studio smoke checks on {platform.platform()} / Python {platform.python_version()}\n")


@check("Python 3.10 or newer, 64-bit")
def _():
    ok = sys.version_info >= (3, 10) and sys.maxsize > 2**32
    return ("PASS" if ok else "FAIL", f"{platform.python_version()} {platform.architecture()[0]}")


for mod, required in (("flask", True), ("PIL", True), ("requests", True), ("pptx", True), ("docx", True), ("waitress", True),
                      ("keyring", True), ("cryptography", True), ("webview", False)):
    @check(f"library '{mod}' installed")
    def _(mod=mod, required=required):
        try:
            m = importlib.import_module(mod)
        except ImportError:
            return ("FAIL" if required else "WARN", "missing" + ("" if required else " - desktop window will fall back to the browser"))
        return getattr(m, "__version__", getattr(m, "VERSION", "ok"))


@check("Pillow can draw text (FreeType) and was built with PNG/JPEG")
def _():
    from PIL import features
    need = {"freetype2": features.check("freetype2"), "zlib": features.check("zlib"), "jpg": features.check("jpg")}
    return ("PASS", "ok") if all(need.values()) else ("FAIL", str(need))


@check("bundled fonts are present")
def _():
    fonts = list((SHARED / "assets" / "fonts").glob("*.ttf"))
    return ("PASS", f"{len(fonts)} fonts") if len(fonts) >= 10 else ("FAIL", f"only {len(fonts)} found")


@check("data folder is writable (projects, uploads, secrets)")
def _():
    from studio.paths import studio_home
    home = studio_home()
    probe = home / "write-test.tmp"
    probe.write_text("x", encoding="utf-8")
    probe.unlink()
    notes = []
    low = str(home).lower()
    if "onedrive" in low:
        notes.append("home is inside OneDrive: sync can briefly lock files (the app retries saves)")
    if " " in str(home) or not str(home).isascii():
        notes.append("path has spaces or non-ASCII characters (supported)")
    return ("WARN" if "onedrive" in low else "PASS", f"{home}" + (" | " + "; ".join(notes) if notes else ""))


for port, who in ((5000, "website"), (5001, "desktop app sign-in"), (8765, "demo Canva")):
    @check(f"port {port} is free ({who})")
    def _(port=port, who=who):
        with socket.socket() as s:
            s.settimeout(0.3)
            busy = s.connect_ex(("127.0.0.1", port)) == 0
        return ("WARN", "in use - close the other program or change the port") if busy else ("PASS", "free")


@check("Windows: Edge WebView2 runtime installed (needed for the desktop window)")
def _():
    if not IS_WIN:
        return ("SKIP", "not Windows")
    import winreg
    guid = r"{F3017226-FE2A-4295-8BDF-00C3A9A7E4C5}"
    for hive, path in ((winreg.HKEY_LOCAL_MACHINE, rf"SOFTWARE\WOW6432Node\Microsoft\EdgeUpdate\Clients\{guid}"),
                       (winreg.HKEY_LOCAL_MACHINE, rf"SOFTWARE\Microsoft\EdgeUpdate\Clients\{guid}"),
                       (winreg.HKEY_CURRENT_USER, rf"Software\Microsoft\EdgeUpdate\Clients\{guid}")):
        try:
            with winreg.OpenKey(hive, path) as k:
                return ("PASS", "version " + winreg.QueryValueEx(k, "pv")[0])
        except OSError:
            continue
    return ("WARN", "not found - install 'Evergreen WebView2 Runtime' from Microsoft; the app will open in your browser instead")


@check("Windows: Credential Manager holds the encryption key")
def _():
    if not IS_WIN:
        return ("SKIP", "not Windows")
    import keyring
    name = type(keyring.get_keyring()).__name__
    return ("PASS", name) if "WinVault" in name else ("WARN", f"backend is {name}; a key file is used instead")


@check("Windows: long paths enabled")
def _():
    if not IS_WIN:
        return ("SKIP", "not Windows")
    import winreg
    try:
        with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, r"SYSTEM\CurrentControlSet\Control\FileSystem") as k:
            return ("PASS", "enabled") if winreg.QueryValueEx(k, "LongPathsEnabled")[0] else ("WARN", "disabled - keep the unzip folder path short (under ~200 characters)")
    except OSError:
        return ("WARN", "could not read the setting")


@check("Windows: reserved device names never become file names")
def _():
    from studio.models import slug
    bad = [n for n in ("CON", "PRN", "AUX", "NUL", "COM1", "LPT1") if not slug(n).startswith("design-")]
    return ("FAIL", f"not protected: {bad}") if bad else ("PASS", "ok")


@check("Windows: unzip location length")
def _():
    n = len(str(SHARED))
    return ("WARN", f"{n} characters - shorter than ~200 is safer on Windows") if n > 200 else ("PASS", f"{n} characters")


@check("pipeline: render, export .pptx/.docx and re-open them")
def _():
    import io

    from docx import Document
    from pptx import Presentation

    from studio import generator as G
    from studio import office
    from studio import render as R
    deck, doc = G.project_from_example("yoga-presentation"), G.project_from_example("workout-blog-refined")
    assert R.render_pages(deck, 0.3) and R.render_pages(doc, 0.3)
    assert len(Presentation(io.BytesIO(office.presentation_to_pptx(deck))).slides) == 6
    assert Document(io.BytesIO(office.document_to_docx(doc))).paragraphs[0].text == doc["title"]
    return "ok"


@check("server: starts on a free port and serves the app")
def _():
    from studio.desktop import make_server
    from studio.webapp import create_app
    with tempfile.TemporaryDirectory() as home:
        with socket.socket() as s:
            s.bind(("127.0.0.1", 0))
            port = s.getsockname()[1]
        server = make_server(create_app(home=Path(home)), "127.0.0.1", port)
        t = threading.Thread(target=server.run, daemon=True)
        t.start()
        try:
            body = b""
            for _ in range(60):
                try:
                    body = urllib.request.urlopen(f"http://127.0.0.1:{port}/api/bootstrap", timeout=2).read()
                    break
                except OSError:
                    time.sleep(0.1)
            assert b"examples" in body
            page = urllib.request.urlopen(f"http://127.0.0.1:{port}/static/app.js", timeout=2).read()
            assert len(page) > 10000
        finally:
            server.close()
            t.join(timeout=5)
    return "ok"


print()
counts = {k: sum(r["status"] == k for r in results) for k in ("PASS", "WARN", "FAIL", "SKIP")}
print("Summary: " + ", ".join(f"{v} {k}" for k, v in counts.items()))
if "--json" in sys.argv:
    out = Path(sys.argv[sys.argv.index("--json") + 1])
    out.write_text(json.dumps({"platform": platform.platform(), "python": platform.python_version(), "results": results}, indent=2), encoding="utf-8")
sys.exit(1 if counts["FAIL"] else 0)
