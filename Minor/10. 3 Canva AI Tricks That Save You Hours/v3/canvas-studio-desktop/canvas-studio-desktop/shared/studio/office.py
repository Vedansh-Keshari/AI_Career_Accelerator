"""Editable Office files (.pptx, .docx) built from a project.

Canva imports these with live text boxes, shapes and pictures, so what is pushed can be
edited in Canva, unlike a PDF of page images. Slide backgrounds and decorations are placed
as one picture behind native text; everything you can drag in the editor is a native object.
"""
from __future__ import annotations

import io

from PIL import Image, ImageOps

from . import options as O
from .elements import default_elements
from .render import Look, fit_element_text, render_background
from .render.base import resolve_color
from .render.free import parse_paragraphs

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

# Fonts that exist on typical machines and in Canva; the bundled preview fonts are not installed everywhere.
OFFICE_FONTS = {"sans": "Poppins", "serif": "Georgia", "dsans": "Arial", "mono": "Courier New"}
SLIDE_SHORT_SIDE_IN = 7.5


def _rgb(hex_color: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def presentation_to_pptx(project: dict) -> bytes:
    """One native slide per slide: styled background picture + editable text, shapes and images."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    look = Look(project)
    bw, bh = O.FORMATS[project["options"]["format"]]
    scale_in = SLIDE_SHORT_SIDE_IN / min(bw, bh)
    w_in, h_in = bw * scale_in, bh * scale_in
    emu = 914400
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(w_in * emu)), Emu(int(h_in * emu))
    blank = prs.slide_layouts[6]
    u = min(bw, bh) / 1080  # pixels per design unit at full size
    font_name = OFFICE_FONTS.get(look.kind, "Arial")

    for i, slide_data in enumerate(project["slides"]):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(look.bg)
        bg = render_background(project, i, (bw // 2, bh // 2), look)
        buf = io.BytesIO()
        bg.save(buf, "PNG")
        buf.seek(0)
        slide.shapes.add_picture(buf, 0, 0, prs.slide_width, prs.slide_height)
        elements = slide_data["elements"] if slide_data.get("free") else default_elements(project, i)
        for el in elements:
            left, top = Emu(int(el["x"] * w_in * emu)), Emu(int(el["y"] * h_in * emu))
            width, height = Emu(int(el["w"] * w_in * emu)), Emu(int(el["h"] * h_in * emu))
            if el["type"] in ("rect", "ellipse"):
                kind = MSO_SHAPE.OVAL if el["type"] == "ellipse" else (MSO_SHAPE.ROUNDED_RECTANGLE if el["radius"] > 0 else MSO_SHAPE.RECTANGLE)
                shape = slide.shapes.add_shape(kind, left, top, width, height)
                if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                    shape.adjustments[0] = min(0.5, el["radius"])
                shape.fill.solid()
                shape.fill.fore_color.rgb = _rgb(resolve_color(look, el["fill"]))
                shape.line.fill.background()
                shape.shadow.inherit = False
                if el["opacity"] < 1:
                    clr = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(el["opacity"] * 100000))}))
            elif el["type"] == "image":
                try:
                    photo = Image.open(el["image"]).convert("RGB")
                except (OSError, ValueError):
                    continue
                fitted = ImageOps.fit(photo, (max(2, int(el["w"] * bw)), max(2, int(el["h"] * bh))), Image.Resampling.LANCZOS)
                pbuf = io.BytesIO()
                fitted.save(pbuf, "JPEG", quality=90)
                pbuf.seek(0)
                slide.shapes.add_picture(pbuf, left, top, width, height)
            else:
                box = slide.shapes.add_textbox(left, top, width, height)
                tf = box.text_frame
                tf.word_wrap = True
                tf.auto_size = None
                pad = Emu(int(0.1 * emu)) if el.get("fill") else Emu(0)
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = pad
                tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[el["valign"]]
                if el.get("fill"):
                    box.fill.solid()
                    box.fill.fore_color.rgb = _rgb(resolve_color(look, el["fill"]))
                fit = fit_element_text(el, look, el["w"] * bw, el["h"] * bh, u)
                pt = max(6.0, fit["size"] / u * 0.5)  # 1080 design units = 7.5 in = 540 pt
                color = _rgb(resolve_color(look, el["color"]))
                align = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[el["align"]]
                for n, (bullet, text) in enumerate(parse_paragraphs(el["text"])):
                    para = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
                    para.alignment = align
                    para.line_spacing = 1.1
                    run = para.add_run()
                    run.text = ("\u2022 " if bullet else "") + text
                    run.font.size = Pt(pt if text else pt * 0.5)
                    run.font.bold = bool(el["bold"])
                    run.font.name = font_name
                    run.font.color.rgb = color
    prs.core_properties.title = project["title"][:120]
    prs.core_properties.author = "Canvas Studio"
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def document_to_docx(project: dict) -> bytes:
    """A Word document with a real title, headings, paragraphs and bullet lists."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    look = Look(project)
    font_name = OFFICE_FONTS.get(look.kind, "Arial")
    pal = look.pal
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)  # A4
    section.left_margin = section.right_margin = Inches(0.95)
    section.top_margin = section.bottom_margin = Inches(0.9)

    def style_font(style_name: str, size: float, bold: bool = False, color: str | None = None):
        st = doc.styles[style_name]
        st.font.name = font_name
        st.font.size = Pt(size)
        st.font.bold = bold
        if color:
            st.font.color.rgb = RGBColor.from_string(color.lstrip("#").upper())
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(__import__("docx.oxml.ns", fromlist=["qn"]).qn(attr), font_name)

    heading_color = look.heading if look.mode != "primary" else pal["primary"]
    style_font("Normal", 11, color=pal["ink"])
    style_font("Title", 30, True, pal["primary"])
    style_font("Heading 1", 16, True, heading_color)
    style_font("List Bullet", 11, color=pal["ink"])
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.25

    doc.add_paragraph(project["title"], style="Title")
    for sec in project["doc"]["sections"]:
        if sec["heading"].strip():
            doc.add_paragraph(sec["heading"].strip(), style="Heading 1")
        for bullet, text in parse_paragraphs(sec["body"]):
            if not text:
                continue
            doc.add_paragraph(text, style="List Bullet" if bullet else "Normal")
    doc.core_properties.title = project["title"][:120]
    doc.core_properties.author = "Canvas Studio"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
